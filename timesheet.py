import io
import calendar
import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.pdfgen import canvas
from engine import get_shift_for_crew, get_holiday_type

MESI_MAIUSCOLO = [
    "", "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
    "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE"
]

GIORNI_SETT_LETTERE = ["L", "M", "M", "G", "V", "S", "D"]

def get_month_crew_data(year: int, month: int, crew_num: int) -> dict:
    """
    Calcola per ciascun giorno del mese (1..31) i valori teorici standard del turno.
    """
    _, num_days = calendar.monthrange(year, month)
    days_data = {}

    tot_notturna = 0
    tot_navigazione = 0
    tot_festivo = 0
    tot_buoni_pasto = 0

    for day in range(1, 32):
        if day > num_days:
            days_data[day] = {
                "exists": False
            }
            continue

        target_date = datetime.date(year, month, day)
        shift = get_shift_for_crew(crew_num, target_date)
        is_hol, _, _ = get_holiday_type(target_date)
        stato = shift.get("stato", "")
        giorno_lett = GIORNI_SETT_LETTERE[target_date.weekday()]

        eq_str = ""
        is_riserva = False
        if stato.startswith("R"):
            is_riserva = True
            eq_str = "R"
        elif stato in ["20", "08", "08:20"]:
            eq_str = str(crew_num)

        dalle_str = ""
        alle_str = ""
        v_nott = None
        v_nav = None
        v_fest = None
        v_buono = None

        if stato == "20":
            dalle_str = "20"
            alle_str = "24"
            v_nott = 4
            v_nav = 4
            v_buono = 1
            if is_hol:
                v_fest = 1
        elif stato == "08":
            dalle_str = "00"
            alle_str = "08"
            v_nott = 6
            v_nav = 8
            v_buono = 1
            if is_hol:
                v_fest = 1
        elif stato == "08:20":
            dalle_str = "08"
            alle_str = "20"
            v_nav = 12
            v_buono = 1
            if is_hol:
                v_fest = 1

        if v_nott:
            tot_notturna += v_nott
        if v_nav:
            tot_navigazione += v_nav
        if v_fest:
            tot_festivo += v_fest
        if v_buono:
            tot_buoni_pasto += v_buono

        days_data[day] = {
            "exists": True,
            "giorno_lett": giorno_lett,
            "is_holiday": is_hol,
            "is_riserva": is_riserva,
            "equipaggio": eq_str,
            "dalle": dalle_str,
            "alle": alle_str,
            "notturna": v_nott,
            "navigazione": v_nav,
            "festivo": v_fest,
            "buoni_pasto": v_buono
        }

    return {
        "year": year,
        "month": month,
        "crew_num": crew_num,
        "num_days": num_days,
        "days": days_data,
        "tot_notturna": tot_notturna,
        "tot_navigazione": tot_navigazione,
        "tot_festivo": tot_festivo,
        "tot_buoni_pasto": tot_buoni_pasto
    }


def generate_monthly_timesheet_pdf(year: int, month: int, crew_num: int) -> bytes:
    """
    Genera il foglio presenze mensile da zero in PDF vettoriale A4 Landscape nativo (842 x 595 pt),
    riproducendo esattamente la grafica del modulo aziendale Tripmare.
    """
    PAGE_W = 841.89
    PAGE_H = 595.28

    data = get_month_crew_data(year, month, crew_num)
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=(PAGE_W, PAGE_H))

    # Margini del foglio A4
    MARGIN_LEFT = 28.0
    MARGIN_RIGHT = PAGE_W - 28.0
    MARGIN_TOP = PAGE_H - 24.0

    # Dimensionamento colonne
    LABEL_COL_W = 168.0
    TOTAL_COL_W = 26.0
    GRID_WIDTH = MARGIN_RIGHT - MARGIN_LEFT - LABEL_COL_W - TOTAL_COL_W
    DAY_COL_W = GRID_WIDTH / 31.0

    X_DAYS_START = MARGIN_LEFT + LABEL_COL_W
    X_TOTAL = MARGIN_RIGHT - TOTAL_COL_W

    def get_cx(d: int) -> float:
        return X_DAYS_START + (d - 1) * DAY_COL_W + (DAY_COL_W / 2.0)

    # 1. TESTATA SUPERIORE (Mese, Anno, Cognome, Nome, Codice)
    Y_HEADER_TOP = MARGIN_TOP
    Y_HEADER_BOTTOM = MARGIN_TOP - 30.0

    c.setLineWidth(1.0)
    c.rect(MARGIN_LEFT, Y_HEADER_BOTTOM, MARGIN_RIGHT - MARGIN_LEFT, 30.0, stroke=1, fill=0)

    # Separatori verticali testata
    x_h1 = MARGIN_LEFT + 110.0
    x_h2 = x_h1 + 240.0
    x_h3 = x_h2 + 220.0

    c.setLineWidth(0.6)
    c.line(x_h1, Y_HEADER_BOTTOM, x_h1, Y_HEADER_TOP)
    c.line(x_h2, Y_HEADER_BOTTOM, x_h2, Y_HEADER_TOP)
    c.line(x_h3, Y_HEADER_BOTTOM, x_h3, Y_HEADER_TOP)

    # Testi fissi testata
    c.setFont("Helvetica-Bold", 8.0)
    c.drawString(MARGIN_LEFT + 6.0, Y_HEADER_TOP - 10.0, "Mese")
    c.drawString(x_h1 + 6.0, Y_HEADER_TOP - 10.0, "Cognome")
    c.drawString(x_h2 + 6.0, Y_HEADER_TOP - 10.0, "Nome")
    c.drawString(x_h3 + 6.0, Y_HEADER_TOP - 10.0, "Codice")

    # Mese e Anno compilati
    c.setFont("Helvetica-Bold", 11.0)
    c.drawString(MARGIN_LEFT + 6.0, Y_HEADER_TOP - 23.0, MESI_MAIUSCOLO[month])
    c.drawString(MARGIN_LEFT + 78.0, Y_HEADER_TOP - 23.0, str(year))

    # 2. DEFINIZIONE RIGHE DELLA GRIGLIA PRINCIPALE
    # Lista righe: (tipo_riga, label_riga, altezza_pt)
    rows_def = [
        ("NUMERI_GIORNO", "", 13.0),
        ("GIORNO_SETT", "giorno della settimana", 13.0),
        ("EQUIPAGGIO", "Equipaggio", 13.0),
        ("DALLE", "Orario di servizio   Dalle", 12.0),
        ("ALLE", "                              Alle", 12.0),
        ("STR_1", "1  Str.  Diu. Feriale", 12.0),
        ("STR_2", "2         Nott. Fer./ Diu. Festivo", 12.0),
        ("STR_3", "3         Nott. Festivo", 12.0),
        ("VOCE_4", "4  Servizi Prolungati", 12.0),
        ("VOCE_5", "5  Maggiorazione Notturna", 12.0),
        ("VOCE_6", "6  Indennità Navigazione", 12.0),
        ("VOCE_7", "7  Indennità Sostitutiva", 12.0),
        ("VOCE_8", "8  Festivo", 12.0),
        ("VOCE_9", "9  Lavori non rientranti", 12.0),
        ("VOCE_10", "10 Passaggio", 12.0),
        ("VOCE_11", "11 Risc. Diurno", 12.0),
        ("VOCE_12", "12 Risc. Notturno", 12.0),
        ("VOCE_13", "13 Compensativo/i Maturato/i", 12.0),
        ("VOCE_14", "14 Reperibilità (solo 8-20 00-08)", 12.0),
        ("VOCE_15", "15 Buoni pasto", 12.0),
        ("VOCE_16", "16 Ore Sciopero", 12.0),
        ("VOCE_17", "17", 12.0),
        ("CAPO_COG", "Cognome del Capo Servizio", 15.0),
        ("CAPO_FIRMA", "Firma del Capo Servizio", 28.0),
    ]

    # Calcolo coordinate Y delle righe
    y_current = Y_HEADER_BOTTOM - 4.0
    row_coords = []
    for r_type, r_label, r_h in rows_def:
        y_top = y_current
        y_bot = y_current - r_h
        row_coords.append((r_type, r_label, y_top, y_bot))
        y_current = y_bot

    grid_top = row_coords[0][2]
    grid_bot = row_coords[-1][3]

    # Tracciamento righe orizzontali
    c.setLineWidth(0.5)
    for idx, (r_type, r_label, y_top, y_bot) in enumerate(row_coords):
        if idx == 0:
            c.setLineWidth(1.0)
            c.line(MARGIN_LEFT, y_top, MARGIN_RIGHT, y_top)
            c.setLineWidth(0.5)
        
        # Bordo inferiore riga
        is_thick_line = r_type in ["ALLE", "VOCE_17", "CAPO_FIRMA"]
        c.setLineWidth(0.9 if is_thick_line else 0.5)
        c.line(MARGIN_LEFT, y_bot, MARGIN_RIGHT, y_bot)

        # Scrittura label a sinistra
        c.setFont("Helvetica", 6.8 if "Reperibilità" in r_label else 7.2)
        y_text = y_bot + (y_top - y_bot) / 2.0 - 2.5
        c.drawString(MARGIN_LEFT + 3.0, y_text, r_label)

    # Scrittura "Tot." sopra la colonna totale
    y_top_tot = row_coords[4][2] # Riga Alle
    c.setFont("Helvetica-Bold", 7.5)
    c.drawCentredString(X_TOTAL + TOTAL_COL_W / 2.0, y_top_tot - 9.0, "Tot.")

    # Tracciamento linee verticali delle colonne
    c.setLineWidth(1.0)
    c.line(MARGIN_LEFT, grid_top, MARGIN_LEFT, grid_bot)
    c.line(X_DAYS_START, grid_top, X_DAYS_START, grid_bot)
    c.line(X_TOTAL, grid_top, X_TOTAL, grid_bot)
    c.line(MARGIN_RIGHT, grid_top, MARGIN_RIGHT, grid_bot)

    # Linee interne delle 31 colonne
    c.setLineWidth(0.4)
    for day in range(1, 31):
        x_col = X_DAYS_START + day * DAY_COL_W
        c.line(x_col, grid_top, x_col, grid_bot)

    # 3. POPOLAMENTO DATI DELLA TABELLA
    days_dict = data["days"]

    for day in range(1, 32):
        d_info = days_dict[day]
        cx = get_cx(day)

        if not d_info["exists"]:
            # Traccia linea diagonale nera sui giorni non presenti
            c.setLineWidth(0.8)
            col_x1 = X_DAYS_START + (day - 1) * DAY_COL_W
            col_x2 = col_x1 + DAY_COL_W
            c.line(col_x1, grid_top, col_x2, grid_bot)
            continue

        for r_type, _, y_top, y_bot in row_coords:
            y_mid = y_bot + (y_top - y_bot) / 2.0 - 2.5

            if r_type == "NUMERI_GIORNO":
                c.setFont("Helvetica-Bold", 7.5)
                c.drawCentredString(cx, y_mid, str(day))

            elif r_type == "GIORNO_SETT":
                c.setFont("Helvetica-Bold", 7.5)
                c.drawCentredString(cx, y_mid, d_info["giorno_lett"])
                if d_info["is_holiday"]:
                    c.setLineWidth(0.6)
                    c.ellipse(cx - 5.5, y_mid - 2.5, cx + 5.5, y_mid + 8.5)

            elif r_type == "EQUIPAGGIO":
                if d_info["is_riserva"]:
                    c.setFont("Helvetica-Bold", 5.5)
                    c.drawString(cx - 6.5, y_mid + 2.5, "R")
                elif d_info["equipaggio"]:
                    c.setFont("Helvetica-Bold", 7.5)
                    c.drawCentredString(cx, y_mid, d_info["equipaggio"])

            elif r_type == "DALLE":
                if d_info["dalle"]:
                    c.setFont("Helvetica", 7.0)
                    c.drawCentredString(cx, y_mid, d_info["dalle"])

            elif r_type == "ALLE":
                if d_info["alle"]:
                    c.setFont("Helvetica", 7.0)
                    c.drawCentredString(cx, y_mid, d_info["alle"])

            elif r_type == "VOCE_5":
                if d_info["notturna"]:
                    c.setFont("Helvetica", 7.5)
                    c.drawCentredString(cx, y_mid, str(d_info["notturna"]))

            elif r_type == "VOCE_6":
                if d_info["navigazione"]:
                    c.setFont("Helvetica", 7.5)
                    c.drawCentredString(cx, y_mid, str(d_info["navigazione"]))

            elif r_type == "VOCE_8":
                if d_info["festivo"]:
                    c.setFont("Helvetica", 7.5)
                    c.drawCentredString(cx, y_mid, str(d_info["festivo"]))

            elif r_type == "VOCE_15":
                if d_info["buoni_pasto"]:
                    c.setFont("Helvetica", 7.5)
                    c.drawCentredString(cx, y_mid, str(d_info["buoni_pasto"]))

    # 4. TOTALI COLONNA A DESTRA
    cx_tot = X_TOTAL + TOTAL_COL_W / 2.0
    c.setFont("Helvetica-Bold", 7.5)
    for r_type, _, y_top, y_bot in row_coords:
        y_mid = y_bot + (y_top - y_bot) / 2.0 - 2.5
        if r_type == "VOCE_5" and data["tot_notturna"] > 0:
            c.drawCentredString(cx_tot, y_mid, str(data["tot_notturna"]))
        elif r_type == "VOCE_6" and data["tot_navigazione"] > 0:
            c.drawCentredString(cx_tot, y_mid, str(data["tot_navigazione"]))
        elif r_type == "VOCE_8" and data["tot_festivo"] > 0:
            c.drawCentredString(cx_tot, y_mid, str(data["tot_festivo"]))
        elif r_type == "VOCE_15" and data["tot_buoni_pasto"] > 0:
            c.drawCentredString(cx_tot, y_mid, str(data["tot_buoni_pasto"]))

    # 5. TABELLA RIEPILOGO FERIE E COMPENSATIVI (In basso a sinistra)
    Y_SUMMARY_TOP = grid_bot - 8.0
    SUMMARY_W = 420.0
    c.setLineWidth(0.6)
    c.rect(MARGIN_LEFT, Y_SUMMARY_TOP - 36.0, SUMMARY_W, 36.0, stroke=1, fill=0)

    # Righe riepilogo
    c.line(MARGIN_LEFT, Y_SUMMARY_TOP - 12.0, MARGIN_LEFT + SUMMARY_W, Y_SUMMARY_TOP - 12.0)
    c.line(MARGIN_LEFT, Y_SUMMARY_TOP - 24.0, MARGIN_LEFT + SUMMARY_W, Y_SUMMARY_TOP - 24.0)

    # Colonne riepilogo: Etichetta (100pt), Inizio mese (80pt), Fruito (80pt), Maturato (80pt), Fine mese (80pt)
    c.line(MARGIN_LEFT + 100.0, Y_SUMMARY_TOP, MARGIN_LEFT + 100.0, Y_SUMMARY_TOP - 36.0)
    c.line(MARGIN_LEFT + 180.0, Y_SUMMARY_TOP, MARGIN_LEFT + 180.0, Y_SUMMARY_TOP - 36.0)
    c.line(MARGIN_LEFT + 260.0, Y_SUMMARY_TOP, MARGIN_LEFT + 260.0, Y_SUMMARY_TOP - 36.0)
    c.line(MARGIN_LEFT + 340.0, Y_SUMMARY_TOP, MARGIN_LEFT + 340.0, Y_SUMMARY_TOP - 36.0)

    c.setFont("Helvetica-Bold", 7.0)
    c.drawCentredString(MARGIN_LEFT + 140.0, Y_SUMMARY_TOP - 9.0, "Saldo inizio mese")
    c.drawCentredString(MARGIN_LEFT + 220.0, Y_SUMMARY_TOP - 9.0, "Fruito")
    c.drawCentredString(MARGIN_LEFT + 300.0, Y_SUMMARY_TOP - 9.0, "Maturato")
    c.drawCentredString(MARGIN_LEFT + 380.0, Y_SUMMARY_TOP - 9.0, "Saldo fine mese")

    c.setFont("Helvetica", 7.5)
    c.drawString(MARGIN_LEFT + 6.0, Y_SUMMARY_TOP - 21.0, "Ferie")
    c.drawString(MARGIN_LEFT + 6.0, Y_SUMMARY_TOP - 33.0, "Compensativi")

    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def generate_monthly_timesheet_excel(year: int, month: int, crew_num: int) -> bytes:
    """
    Genera il file Excel (.xlsx) nativo basato sulla struttura del foglio presenze Tripmare.
    """
    data = get_month_crew_data(year, month, crew_num)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Presenze Eq {crew_num}"

    # Stili
    font_bold = Font(name="Calibri", size=9, bold=True)
    font_regular = Font(name="Calibri", size=9)
    font_title = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_border_side = Side(border_style="thin", color="000000")
    double_border_side = Side(border_style="medium", color="000000")
    cell_border = Border(top=thin_border_side, bottom=thin_border_side, left=thin_border_side, right=thin_border_side)

    # 1. Testata (Righe 1 e 2)
    ws.merge_cells("A1:D1")
    ws["A1"] = "Mese"
    ws["A1"].font = font_bold

    ws.merge_cells("A2:C2")
    ws["A2"] = MESI_MAIUSCOLO[month]
    ws["A2"].font = font_title

    ws["D2"] = year
    ws["D2"].font = font_title
    ws["D2"].alignment = align_center

    ws.merge_cells("E1:P1")
    ws["E1"] = "Cognome"
    ws["E1"].font = font_bold

    ws.merge_cells("Q1:AA1")
    ws["Q1"] = "Nome"
    ws["Q1"].font = font_bold

    ws.merge_cells("AB1:AI1")
    ws["AB1"] = "Codice"
    ws["AB1"].font = font_bold

    # Riga 2 (Numeri da 1 a 31)
    for day in range(1, 32):
        col_letter = openpyxl.utils.get_column_letter(4 + day)
        cell = ws[f"{col_letter}2"]
        cell.value = day
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = cell_border

    # 2. Voci Righe
    row_labels = [
        (3, "giorno della settimana"),
        (4, "Equipaggio"),
        (5, "Orario di servizio - Dalle"),
        (6, "Orario di servizio - Alle"),
        (7, "1. Str. Diu. Feriale"),
        (8, "2. Str. Nott. Fer./ Diu. Festivo"),
        (9, "3. Str. Nott. Festivo"),
        (10, "4. Servizi Prolungati"),
        (11, "5. Maggiorazione Notturna"),
        (12, "6. Indennità Navigazione"),
        (13, "7. Indennità Sostitutiva"),
        (14, "8. Festivo"),
        (15, "9. Lavori non rientranti"),
        (16, "10. Passaggio"),
        (17, "11. Risc. Diurno"),
        (18, "12. Risc. Notturno"),
        (19, "13. Compensativo/i Maturato/i"),
        (20, "14. Reperibilità (solo 8-20 00-08)"),
        (21, "15. Buoni pasto"),
        (22, "16. Ore Sciopero"),
        (23, "17."),
        (24, "Cognome del Capo Servizio"),
        (25, "Firma del Capo Servizio")
    ]

    for r_idx, label in row_labels:
        ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=4)
        c_lbl = ws.cell(row=r_idx, column=1)
        c_lbl.value = label
        c_lbl.font = font_bold if r_idx in [3, 4] else font_regular
        c_lbl.alignment = align_left

        for c_i in range(1, 5):
            ws.cell(row=r_idx, column=c_i).border = cell_border

    # Intestazione Totale in Colonna AJ (36)
    ws.cell(row=6, column=36, value="Tot.").font = font_bold
    ws.cell(row=6, column=36).alignment = align_center

    # 3. Inserimento dati delle 31 giornate
    days_dict = data["days"]
    for day in range(1, 32):
        col_idx = 4 + day
        d_info = days_dict[day]

        for r_idx in range(3, 26):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = cell_border
            cell.alignment = align_center

            if not d_info["exists"]:
                cell.value = "-"
                continue

            if r_idx == 3:
                cell.value = d_info["giorno_lett"]
                cell.font = font_bold
            elif r_idx == 4:
                cell.value = d_info["equipaggio"]
                cell.font = font_bold
            elif r_idx == 5:
                cell.value = d_info["dalle"]
            elif r_idx == 6:
                cell.value = d_info["alle"]
            elif r_idx == 11:
                cell.value = d_info["notturna"]
            elif r_idx == 12:
                cell.value = d_info["navigazione"]
            elif r_idx == 14:
                cell.value = d_info["festivo"]
            elif r_idx == 21:
                cell.value = d_info["buoni_pasto"]

    # 4. Totali nella colonna 36 (AJ)
    tot_map = {
        11: data["tot_notturna"],
        12: data["tot_navigazione"],
        14: data["tot_festivo"],
        21: data["tot_buoni_pasto"]
    }

    for r_idx in range(7, 24):
        cell_tot = ws.cell(row=r_idx, column=36)
        cell_tot.border = cell_border
        cell_tot.alignment = align_center
        cell_tot.font = font_bold
        if r_idx in tot_map and tot_map[r_idx] > 0:
            cell_tot.value = tot_map[r_idx]

    # Dimensionamento colonne
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 6
    for day in range(1, 32):
        col_letter = openpyxl.utils.get_column_letter(4 + day)
        ws.column_dimensions[col_letter].width = 4.0
    ws.column_dimensions["AJ"].width = 6.0

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
